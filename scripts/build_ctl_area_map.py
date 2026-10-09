"""Build public/cable-ctl-area-map.json from the Control Cable Schedule.

The schedule states each cable's building number in SWGR LOCATION (FROM) and
LOAD LOCATION (TO), e.g. "1.2_LEB FOR BLOCK #1" or "6_BLOCK #1 GTG #1 UAT".
Control FROM/TO tags often carry suffixes ("B1-PRP-64701 IED-1"), so a tag-keyed
lookup misses them; keying by cable number takes the schedule's own answer.

Usage: python scripts/build_ctl_area_map.py "<Control Cable Schedule .xlsx>"
Rerun whenever a new Control Cable Schedule revision is applied.
"""
import io
import json
import re
import sys

from openpyxl import load_workbook

SHEET = 'Control Cable Schedule_공유'
COL_NO, COL_FROM_LOC, COL_TO_LOC = 9, 17, 20  # 0-based: NUMBER, SWGR LOCATION, LOAD LOCATION
OUT = 'public/cable-ctl-area-map.json'


def area(loc, cable_no):
    m = re.match(r'\s*(\d+(?:\.\d+)?)\s*_', str(loc or ''))
    if not m:
        return None
    code = m.group(1)
    if code != '1.1':
        return code
    # The Location filter splits the Main Building by block.
    text = str(loc).upper()
    if re.search(r'-2[12]\b|ST-2|ST #2|ST-GCB-2|#2\b', text):
        return '1.1.2'
    if re.search(r'-1[12]\b|ST-1|ST #1|ST-GCB-1|#1\b', text):
        return '1.1.1'
    return '1.1.2' if cable_no.startswith('B2-') else '1.1.1'


def main(path):
    ws = load_workbook(path, read_only=True, data_only=True)[SHEET]
    out = {}
    for r in ws.iter_rows(min_row=9, values_only=True):
        if len(r) <= COL_TO_LOC:
            continue
        n = r[COL_NO]
        if not isinstance(n, str) or not n.strip():
            continue
        n = n.strip()
        f, t = area(r[COL_FROM_LOC], n), area(r[COL_TO_LOC], n)
        # The lookup answers both ends at once, so keep only rows that state both.
        if f and t:
            out[n] = {'from': f, 'to': t}
    with io.open(OUT, 'w', encoding='utf-8', newline='\n') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
        fh.write('\n')
    print(f'{len(out)} cables -> {OUT}')


if __name__ == '__main__':
    main(sys.argv[1])
